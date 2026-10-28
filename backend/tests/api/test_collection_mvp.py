import json
from collections.abc import Generator
from io import BytesIO
from secrets import token_urlsafe
from threading import Lock
from time import monotonic, sleep
from typing import Any
from zipfile import ZipFile

import httpx
import pytest
from cryptography.fernet import Fernet
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.collection.doubao_browser import DoubaoBrowserStatus, DoubaoWebTimeout
from app.collection.normalization import deduplicate_leads, normalize_lead
from app.collection.service import CollectionService
from app.config import Settings
from app.main import create_app

RESULT_FIELDS = [
    "city",
    "company_name",
    "mobile",
    "contact_person_name",
    "contact_person_role",
    "contact_verification_status",
    "contact_evidence_url",
    "contact_evidence_urls",
    "contact_evidence_quote",
    "contact_verified_at",
    "contact_verification_reason",
    "email",
    "wechat",
    "address",
    "service_intro",
    "saved_at",
    "discovery_channel",
    "source_url",
    "matched_keyword",
    "source_confidence",
    "confidence_score",
    "confidence_level",
    "verification_status",
    "verification_models",
    "verification_detail",
    "company_name_verification_status",
    "company_name_evidence_urls",
    "company_name_verification_reason",
    "company_name_verified_at",
    "company_name_verification_attempts",
    "delivery_status",
    "delivery_error",
    "delivered_at",
]


class FakeWebSearchProvider:
    id = "web_search"
    label = "公开网页搜索"
    configured = True

    def list_cities(self) -> list[str]:
        return ["武汉市", "长沙市"]

    def collect(self, *, city: str, keyword: str, request_text: str) -> list[dict[str, Any]]:
        phone_tail = sum(ord(character) for character in f"{city}|{keyword}") % 100_000_000
        return [
            {
                "company_name": f"{city}{keyword}服务中心",
                "mobile": f"138{phone_tail:08d}",
                "email": None,
                "wechat": None,
                "address": f"{city}测试路 1 号",
                "service_intro": keyword,
                "source_url": "https://uri.amap.com/marker?position=114.3,30.5",
            }
        ]


class FailOnceWebSearchProvider(FakeWebSearchProvider):
    def __init__(self) -> None:
        self.failed = False

    def collect(self, *, city: str, keyword: str, request_text: str) -> list[dict[str, Any]]:
        if not self.failed:
            self.failed = True
            from app.collection.providers import ProviderError

            raise ProviderError("temporary provider failure")
        return super().collect(city=city, keyword=keyword, request_text=request_text)


class FakeVerifier:
    id = "deepseek"
    label = "DeepSeek"
    configured = True

    def verify(
        self,
        leads: list[dict[str, Any]],
        *,
        request_text: str,
        city: str,
        keyword: str,
    ) -> list[dict[str, Any]]:
        return [{"relevant": True, "reason": f"{city}{keyword}符合{request_text}"} for _ in leads]


class EmptyWebSearchProvider(FakeWebSearchProvider):
    def collect(self, *, city: str, keyword: str, request_text: str) -> list[dict[str, Any]]:
        return []


class SparseWebSearchProvider(FakeWebSearchProvider):
    def __init__(self) -> None:
        self.visited_cities: list[str] = []

    def collect(self, *, city: str, keyword: str, request_text: str) -> list[dict[str, Any]]:
        self.visited_cities.append(city)
        if city == "武汉市":
            return []
        return super().collect(city=city, keyword=keyword, request_text=request_text)


class ConcurrentProbeProvider(FakeWebSearchProvider):
    def __init__(self) -> None:
        self.lock = Lock()
        self.active = 0
        self.max_active = 0

    def collect(self, *, city: str, keyword: str, request_text: str) -> list[dict[str, Any]]:
        with self.lock:
            self.active += 1
            self.max_active = max(self.max_active, self.active)
        try:
            sleep(0.05)
            return super().collect(city=city, keyword=keyword, request_text=request_text)
        finally:
            with self.lock:
                self.active -= 1


class MustNotRunVerifier(FakeVerifier):
    def verify(
        self,
        leads: list[dict[str, Any]],
        *,
        request_text: str,
        city: str,
        keyword: str,
    ) -> list[dict[str, Any]]:
        raise AssertionError("model must not run without collected evidence")


class FakeDoubaoWebGateway:
    def __init__(self) -> None:
        self.start_calls = 0
        self.reconnect_calls = 0

    def status(self) -> DoubaoBrowserStatus:
        return DoubaoBrowserStatus("waiting_for_login", "需要登录豆包网页", None)

    def start(self) -> DoubaoBrowserStatus:
        self.start_calls += 1
        return self.status()

    def reconnect(self) -> DoubaoBrowserStatus:
        self.reconnect_calls += 1
        return DoubaoBrowserStatus("ready", "豆包网页已重新连接", None)

    def shutdown(self) -> None:
        return None


class TimeoutDoubaoWebGateway(FakeDoubaoWebGateway):
    def start(self) -> DoubaoBrowserStatus:
        raise DoubaoWebTimeout("启动豆包网页浏览器超时")


class TimeoutStatusDoubaoWebGateway(FakeDoubaoWebGateway):
    def status(self) -> DoubaoBrowserStatus:
        raise DoubaoWebTimeout("豆包网页状态查询超时")


class FakeDedicatedChromeLauncher:
    def __init__(self, *, running: bool) -> None:
        self.running = running
        self.start_calls = 0

    def is_running(self) -> bool:
        return self.running

    def start(self) -> bool:
        self.start_calls += 1
        started = not self.running
        self.running = True
        return started


class ReenrichmentVerifier(FakeVerifier):
    def __init__(self, model_id: str) -> None:
        self.id = model_id
        self.label = "豆包" if model_id == "doubao" else "DeepSeek"
        self.discovery_calls = 0
        self.enrichment_calls = 0

    def discover_companies(
        self,
        *,
        request_text: str,
        city: str,
        keywords: list[str],
        limit: int = 100,
    ) -> list[dict[str, Any]]:
        self.discovery_calls += 1
        return [{
            "company_name": f"{city}真实软装有限公司",
            "matched_keyword": keywords[0],
            "reason": request_text,
            "source_confidence": 80,
        }]

    def enrich_company(
        self,
        *,
        request_text: str,
        city: str,
        company: dict[str, Any],
        required_contacts: list[str],
    ) -> dict[str, Any]:
        self.enrichment_calls += 1
        result = {
            "company_name": company["company_name"],
            "mobile": "13800138009",
            "email": None,
            "wechat": None,
            "address": f"{city}公开地址",
            "service_intro": request_text,
            "source_url": "https://example.test/company",
            "source_confidence": 82,
        }
        result["mobile"] = "139" + "12345678"
        return result


@pytest.fixture
def collection_client(tmp_path) -> Generator[TestClient, None, None]:
    settings = Settings(
        app_env="test",
        database_url="sqlite+pysqlite:///:memory:",
        app_master_key=Fernet.generate_key().decode("ascii"),
        session_secret=token_urlsafe(32),
        data_dir=tmp_path,
        amap_api_key=token_urlsafe(24),
        erp_auto_push=False,
    )
    client = TestClient(create_app(settings))
    client.app.state.collection_service.providers = {"web_search": FakeWebSearchProvider()}
    client.app.state.collection_service.company_name_verifier = None
    client.app.state.collection_service.contact_verifier = None
    client.app.state.collection_service.verifiers = {"deepseek": FakeVerifier()}
    try:
        yield client
    finally:
        client.app.state.collection_service.shutdown()
        client.close()


def create_task(client: TestClient, cities: list[str], keywords: list[str]) -> dict:
    response = client.post(
        "/api/v1/collection-tasks",
        json={
            "name": "家具 B 端五城试点",
            "request_text": "寻找全案软装、设计和整装公司",
            "cities": cities,
            "keywords": keywords,
            "providers": ["web_search"],
            "models": ["deepseek"],
            "required_contacts": ["mobile"],
        },
    )
    assert response.status_code == 201
    return response.json()


def wait_for_task(client: TestClient, task_id: str, timeout_seconds: float = 2) -> dict:
    deadline = monotonic() + timeout_seconds
    while True:
        response = client.get(f"/api/v1/collection-tasks/{task_id}")
        assert response.status_code == 200
        task = response.json()
        if task["status"] in {"completed", "completed_with_errors", "paused", "terminated"}:
            return task
        if monotonic() >= deadline:
            raise AssertionError(f"task {task_id} did not reach a terminal state")
        sleep(0.01)


def test_create_start_and_status(collection_client: TestClient) -> None:
    task = create_task(collection_client, ["武汉", "长沙"], ["软装公司", "整装公司"])

    assert task["status"] == "draft"
    assert task["total_units"] == 4
    assert task["completed_units"] == 0
    assert set(task["city_progress"]) == {"武汉", "长沙"}


    started = collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    assert started.status_code == 202

    status = wait_for_task(collection_client, task["id"])
    assert status["status"] == "completed"
    assert status["completed_units"] == 4
    assert status["lead_count"] == 4

    repeated = collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    assert repeated.status_code == 202
    assert collection_client.get(f"/api/v1/collection-tasks/{task['id']}").json()["lead_count"] == 4


def test_doubao_web_status_and_start_are_sanitized(collection_client: TestClient) -> None:
    gateway = FakeDoubaoWebGateway()
    launcher = FakeDedicatedChromeLauncher(running=True)
    collection_client.app.state.collection_service.doubao_web_gateway = gateway
    collection_client.app.state.doubao_web_launcher = launcher

    status_response = collection_client.get("/api/v1/collection-tasks/doubao-web/status")
    start_response = collection_client.post("/api/v1/collection-tasks/doubao-web/start")

    expected = {
        "enabled": True,
        "state": "waiting_for_login",
        "detail": "需要登录豆包网页",
        "current_company": None,
    }
    assert status_response.status_code == 200
    assert start_response.status_code == 200
    assert status_response.json() == expected
    assert start_response.json() == expected
    assert launcher.start_calls == 1
    assert gateway.start_calls == 1
    assert "cookie" not in status_response.text.lower()


def test_doubao_web_status_translates_timeout_to_non_error_state(
    collection_client: TestClient,
) -> None:
    collection_client.app.state.collection_service.doubao_web_gateway = TimeoutStatusDoubaoWebGateway()
    collection_client.app.state.doubao_web_launcher = FakeDedicatedChromeLauncher(running=True)

    response = collection_client.get("/api/v1/collection-tasks/doubao-web/status")

    assert response.status_code == 200
    assert response.json() == {
        "enabled": True,
        "state": "timeout",
        "detail": "豆包网页状态查询超时，任务继续运行",
        "current_company": None,
    }


def test_doubao_web_status_reports_stopped_without_launching(
    collection_client: TestClient,
) -> None:
    gateway = FakeDoubaoWebGateway()
    launcher = FakeDedicatedChromeLauncher(running=False)
    collection_client.app.state.collection_service.doubao_web_gateway = gateway
    collection_client.app.state.doubao_web_launcher = launcher

    response = collection_client.get("/api/v1/collection-tasks/doubao-web/status")

    assert response.status_code == 200
    assert response.json()["state"] == "browser_stopped"
    assert launcher.start_calls == 0
    assert gateway.start_calls == 0


def test_doubao_web_start_translates_timeout_to_service_unavailable(
    collection_client: TestClient,
) -> None:
    collection_client.app.state.collection_service.doubao_web_gateway = TimeoutDoubaoWebGateway()
    collection_client.app.state.doubao_web_launcher = FakeDedicatedChromeLauncher(running=True)

    response = collection_client.post("/api/v1/collection-tasks/doubao-web/start")

    assert response.status_code == 503
    assert "超时" in response.json()["detail"]


def test_doubao_web_reconnect_does_not_launch_chrome(
    collection_client: TestClient,
) -> None:
    gateway = FakeDoubaoWebGateway()
    launcher = FakeDedicatedChromeLauncher(running=True)
    collection_client.app.state.collection_service.doubao_web_gateway = gateway
    collection_client.app.state.doubao_web_launcher = launcher

    response = collection_client.post("/api/v1/collection-tasks/doubao-web/reconnect")

    assert response.status_code == 200
    assert response.json()["state"] == "ready"
    assert launcher.start_calls == 0
    assert gateway.reconnect_calls == 1


def test_lead_pages_are_ordered_and_include_pagination_metadata(
    collection_client: TestClient,
) -> None:
    task = create_task(collection_client, ["武汉", "长沙"], ["软装公司", "整装公司"])
    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])
    service = collection_client.app.state.collection_service
    stored = service.get_task(task["id"])
    ordered_names = []
    for index, lead in enumerate(reversed(stored["leads"])):
        lead["saved_at"] = f"2026-07-27 09:00:0{index}"
        ordered_names.append(lead["company_name"])
    stored["leads"].reverse()
    stored["lead_count"] = 999
    for progress in stored["city_progress"].values():
        progress["lead_count"] = 999
    service.store.save(stored)

    response = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/leads?offset=1&limit=2")

    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 4
    assert body["offset"] == 1
    assert body["limit"] == 2
    assert [lead["company_name"] for lead in body["items"]] == ordered_names[1:3]

    city_response = collection_client.get(
        f"/api/v1/collection-tasks/{task['id']}/leads?city=武汉&limit=1"
    )
    assert city_response.status_code == 200
    assert city_response.json()["total"] == 2
    assert len(city_response.json()["items"]) == 1
    assert city_response.json()["items"][0]["city"] == "武汉"


@pytest.mark.parametrize("query", ["offset=-1", "limit=0", "limit=1001"])
def test_lead_pagination_rejects_invalid_bounds(
    collection_client: TestClient,
    query: str,
) -> None:
    task = create_task(collection_client, ["武汉"], ["软装公司"])

    response = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/leads?{query}")

    assert response.status_code == 422


def test_landlines_are_excluded() -> None:
    lead = normalize_lead(
        {
            "company_name": "武汉示例软装",
            "mobile": "027-8888-6666",
            "email": "HELLO@EXAMPLE.COM",
        },
        city="武汉",
        provider="amap",
    )

    assert lead["mobile"] is None
    assert lead["email"] == "hello@example.com"
    assert __import__("re").fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}", lead["saved_at"])


def test_contacts_are_deduplicated() -> None:
    first = normalize_lead(
        {"company_name": "甲公司", "mobile": "+86 138-0013-8000", "wechat": "Company-A"},
        city="武汉",
        provider="amap",
    )
    duplicate_mobile = normalize_lead(
        {"company_name": "乙公司", "mobile": "13800138000", "email": "b@example.com"},
        city="武汉",
        provider="amap",
    )
    duplicate_wechat = normalize_lead(
        {"company_name": "丙公司", "wechat": " company-a "},
        city="武汉",
        provider="amap",
    )

    deduplicated = deduplicate_leads([first, duplicate_mobile, duplicate_wechat])

    assert [lead["company_name"] for lead in deduplicated] == ["甲公司"]
    assert deduplicated[0]["mobile"] == "13800138000"
    assert deduplicated[0]["email"] == "b@example.com"
    assert deduplicated[0]["wechat"] == "Company-A"


@pytest.mark.parametrize("export_format", ["csv", "xlsx"])
def test_export_contains_one_file_per_city(
    collection_client: TestClient,
    export_format: str,
    monkeypatch,
) -> None:
    task = create_task(collection_client, ["武汉", "西安"], ["整案公司"])
    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])
    store = collection_client.app.state.collection_service.store
    original_iter_leads = store.iter_leads
    original_load = store.load
    batch_sizes: list[int] = []
    include_leads_values: list[bool] = []

    def recording_iter_leads(task_id: str, batch_size: int = 500):
        batch_sizes.append(batch_size)
        return original_iter_leads(task_id, batch_size=batch_size)

    def recording_load(task_id: str, *, include_leads: bool = True):
        include_leads_values.append(include_leads)
        return original_load(task_id, include_leads=include_leads)

    monkeypatch.setattr(store, "iter_leads", recording_iter_leads)
    monkeypatch.setattr(store, "load", recording_load)

    response = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/export?format={export_format}")

    assert response.status_code == 200
    assert batch_sizes == [500]
    assert include_leads_values == [False, True]
    assert response.headers["content-type"] == "application/zip"
    export_path = collection_client.app.state.collection_service.store.exports_dir / f"{task['id']}-{export_format}.zip"
    assert export_path.read_bytes() == response.content
    with ZipFile(BytesIO(response.content)) as archive:
        assert archive.namelist() == [f"武汉.{export_format}", f"西安.{export_format}"]
        if export_format == "csv":
            header = archive.read("武汉.csv").decode("utf-8-sig").splitlines()[0].split(",")
        else:
            workbook = load_workbook(BytesIO(archive.read("武汉.xlsx")), read_only=True)
            header = [cell.value for cell in next(workbook.active.iter_rows())]
        assert header == RESULT_FIELDS


def test_all_city_scope_expands_from_built_in_catalog(collection_client: TestClient, monkeypatch) -> None:
    monkeypatch.setattr(
        "app.collection.service.tier3_plus_city_names", lambda: ["武汉市", "长沙市"]
    )
    response = collection_client.post(
        "/api/v1/collection-tasks",
        json={
            "name": "全国真实线索",
            "request_text": "寻找软装和整装公司",
            "cities": [],
            "all_cities": True,
            "keywords": ["软装公司"],
            "providers": ["web_search"],
            "models": ["deepseek"],
            "required_contacts": ["mobile"],
        },
    )

    assert response.status_code == 201
    task = response.json()
    assert task["cities"] == ["武汉市", "长沙市"]
    assert task["total_units"] == 2


def test_source_and_model_lists_expose_selectable_capabilities(collection_client: TestClient) -> None:
    collection_client.app.state.collection_service.verifiers = {"deepseek": FakeVerifier()}
    response = collection_client.get("/api/v1/providers")
    models = collection_client.get("/api/v1/models")

    assert response.status_code == 200
    assert response.json()["items"] == [
        {"id": "web_search", "label": "公开网页搜索", "configured": True}
    ]
    assert models.status_code == 200
    items = models.json()["items"]
    assert items[0] == {"id": "deepseek", "label": "DeepSeek", "configured": True}
    assert [item["id"] for item in items] == [
        "deepseek",
        "doubao",
        "qwen",
        "kimi",
        "openai",
        "claude",
        "gemini",
    ]
    assert all(item["configured"] is False for item in items[1:])


def test_selected_model_verifies_collected_leads(collection_client: TestClient) -> None:
    collection_client.app.state.collection_service.verifiers = {"deepseek": FakeVerifier()}
    response = collection_client.post(
        "/api/v1/collection-tasks",
        json={
            "name": "模型核验任务",
            "request_text": "寻找全案软装公司",
            "cities": ["武汉市"],
            "keywords": ["软装公司"],
            "providers": ["web_search"],
            "models": ["deepseek"],
            "required_contacts": ["mobile"],
        },
    )

    assert response.status_code == 201
    task = response.json()
    assert task["models"] == ["deepseek"]
    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])
    unverified = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/leads").json()["items"][0]
    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/verify")
    lead = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/leads").json()["items"][0]
    assert unverified["verification_status"] == "unverified"
    assert lead["verification_status"] == "single_model_match"
    assert lead["verification_models"] == ["deepseek"]
    assert lead["verification_detail"]["deepseek"]["relevant"] is True


def test_unknown_or_unconfigured_model_is_rejected(collection_client: TestClient) -> None:
    response = collection_client.post(
        "/api/v1/collection-tasks",
        json={
            "name": "错误模型任务",
            "request_text": "寻找全案软装公司",
            "cities": ["武汉市"],
            "keywords": ["软装公司"],
            "providers": ["web_search"],
            "models": ["unknown-model"],
        },
    )

    assert response.status_code == 422


def test_model_is_not_called_when_source_returns_no_leads(collection_client: TestClient) -> None:
    service = collection_client.app.state.collection_service
    service.providers = {"web_search": EmptyWebSearchProvider()}
    service.verifiers = {"deepseek": MustNotRunVerifier()}
    response = collection_client.post(
        "/api/v1/collection-tasks",
        json={
            "name": "空结果任务",
            "request_text": "寻找全案软装公司",
            "cities": ["武汉市"],
            "keywords": ["软装公司"],
            "providers": ["web_search"],
            "models": ["deepseek"],
            "required_contacts": ["mobile"],
        },
    )
    task = response.json()

    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    status_response = wait_for_task(collection_client, task["id"])

    assert status_response["status"] == "completed_with_errors"
    assert status_response["verification_errors"] == 0
    assert status_response["empty_units"] == 1


def test_city_below_target_does_not_block_following_cities(collection_client: TestClient) -> None:
    provider = SparseWebSearchProvider()
    collection_client.app.state.collection_service.providers = {"web_search": provider}
    task = create_task(collection_client, ["武汉市", "长沙市"], ["软装公司"])

    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    status_response = wait_for_task(collection_client, task["id"])

    assert provider.visited_cities == ["武汉市", "长沙市"]
    assert status_response["status"] == "completed"
    assert status_response["lead_count"] == 1
    assert status_response["empty_units"] == 1


def test_collection_service_honors_configured_task_concurrency(collection_client: TestClient) -> None:
    provider = ConcurrentProbeProvider()
    service = collection_client.app.state.collection_service
    service.providers = {"web_search": provider}
    first = create_task(collection_client, ["武汉市"], ["软装公司"])
    second = create_task(collection_client, ["长沙市"], ["软装公司"])

    service.start_task(first["id"])
    service.start_task(second["id"])
    wait_for_task(collection_client, first["id"])
    wait_for_task(collection_client, second["id"])

    assert provider.max_active == collection_client.app.state.settings.collection_max_concurrent_tasks


def test_new_service_resumes_interrupted_running_task(collection_client: TestClient) -> None:
    service = collection_client.app.state.collection_service
    task = create_task(collection_client, ["武汉市"], ["软装公司"])
    stored = service.get_task(task["id"])
    stored["status"] = "running"
    service.store.save(stored)

    recovered = CollectionService(service.store, service.providers, service.verifiers)
    try:
        status = wait_for_task(collection_client, task["id"])
        assert status["status"] == "completed"
        assert status["lead_count"] == 1
    finally:
        recovered.shutdown()


def test_task_requires_at_least_one_configured_model(collection_client: TestClient) -> None:
    response = collection_client.post(
        "/api/v1/collection-tasks",
        json={
            "name": "缺少模型",
            "request_text": "寻找软装公司",
            "cities": ["武汉市"],
            "keywords": ["软装公司"],
            "providers": ["web_search"],
            "models": [],
            "required_contacts": ["mobile"],
        },
    )

    assert response.status_code == 422


def test_all_selected_contacts_are_required(collection_client: TestClient) -> None:
    collection_client.app.state.collection_service.verifiers = {"deepseek": FakeVerifier()}
    response = collection_client.post(
        "/api/v1/collection-tasks",
        json={
            "name": "邮箱和手机都必填",
            "request_text": "寻找软装公司",
            "cities": ["武汉市"],
            "keywords": ["软装公司"],
            "providers": ["web_search"],
            "models": ["deepseek"],
            "required_contacts": ["mobile", "email"],
        },
    )
    task = response.json()

    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])

    assert collection_client.get(f"/api/v1/collection-tasks/{task['id']}/leads").json()["items"] == []


class RejectingVerifier(FakeVerifier):
    def verify(
        self,
        leads: list[dict[str, Any]],
        *,
        request_text: str,
        city: str,
        keyword: str,
    ) -> list[dict[str, Any]]:
        return [{"relevant": False, "reason": "不符合任务"} for _ in leads]


class FailingDoubaoVerifier(FakeVerifier):
    id = "doubao"
    label = "豆包"

    def verify(
        self,
        leads: list[dict[str, Any]],
        *,
        request_text: str,
        city: str,
        keyword: str,
    ) -> list[dict[str, Any]]:
        from app.collection.models import ModelVerificationError

        raise ModelVerificationError("doubao unavailable")


class DiscoveringVerifier(FakeVerifier):
    def discover(self, *, request_text: str, city: str, keyword: str) -> list[dict[str, Any]]:
        return [{
            "company_name": f"{city}{keyword}模型客户",
            "mobile": "13800138000",
            "email": None,
            "wechat": None,
            "address": f"{city}模型路 1 号",
            "service_intro": request_text,
            "source_url": None,
            "matched_keyword": keyword,
            "source_confidence": 76,
        }]


class DiscoverWithoutAutomaticVerification(DiscoveringVerifier):
    def verify(
        self,
        leads: list[dict[str, Any]],
        *,
        request_text: str,
        city: str,
        keyword: str,
    ) -> list[dict[str, Any]]:
        raise AssertionError("candidate collection must not verify automatically")


class DiscoveringCompanyWithoutRequestedContact(DiscoveringVerifier):
    def discover(self, *, request_text: str, city: str, keyword: str) -> list[dict[str, Any]]:
        item = super().discover(request_text=request_text, city=city, keyword=keyword)[0]
        item["mobile"] = "027-88886666"
        return [item]


class DiscoverFailOnceVerifier(DiscoveringVerifier):
    def __init__(self) -> None:
        self.calls = 0

    def discover(self, *, request_text: str, city: str, keyword: str) -> list[dict[str, Any]]:
        self.calls += 1
        if self.calls == 1:
            from app.collection.models import ModelVerificationError

            raise ModelVerificationError("temporary model discovery failure")
        return super().discover(request_text=request_text, city=city, keyword=keyword)


def test_model_candidates_are_shown_before_verification_when_web_is_empty(
    collection_client: TestClient,
) -> None:
    service = collection_client.app.state.collection_service
    service.providers = {"web_search": EmptyWebSearchProvider()}
    service.verifiers = {"deepseek": DiscoverWithoutAutomaticVerification()}
    task = create_task(collection_client, ["武汉市"], ["软装公司"])

    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])
    leads = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/leads").json()["items"]

    assert len(leads) == 1
    assert leads[0]["company_name"] == "武汉市软装公司模型客户"
    assert leads[0]["discovery_channel"] == "deepseek"
    assert leads[0]["verification_status"] == "unverified"
    assert leads[0]["confidence_level"] == "待验证"


def test_model_company_candidate_is_shown_when_requested_contact_is_not_yet_verified(
    collection_client: TestClient,
) -> None:
    service = collection_client.app.state.collection_service
    service.providers = {"web_search": EmptyWebSearchProvider()}
    service.verifiers = {"deepseek": DiscoveringCompanyWithoutRequestedContact()}
    task = create_task(collection_client, ["武汉市"], ["软装公司"])

    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])
    leads = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/leads").json()["items"]

    assert len(leads) == 1
    assert leads[0]["mobile"] is None
    assert leads[0]["verification_status"] == "unverified"


def test_model_candidate_discovery_retries_once(collection_client: TestClient) -> None:
    service = collection_client.app.state.collection_service
    verifier = DiscoverFailOnceVerifier()
    service.providers = {"web_search": EmptyWebSearchProvider()}
    service.verifiers = {"deepseek": verifier}
    task = create_task(collection_client, ["武汉市"], ["软装公司"])

    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])
    leads = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/leads").json()["items"]

    assert verifier.calls == 2
    assert len(leads) == 1


def test_model_only_task_is_rejected_without_public_web_source(collection_client: TestClient) -> None:
    collection_client.app.state.collection_service.verifiers = {"deepseek": DiscoveringVerifier()}
    response = collection_client.post(
        "/api/v1/collection-tasks",
        json={
            "name": "大模型直接找客户",
            "request_text": "寻找软装公司",
            "cities": ["武汉市"],
            "keywords": ["软装公司"],
            "providers": [],
            "models": ["deepseek"],
            "required_contacts": ["mobile"],
        },
    )

    assert response.status_code == 422


def test_model_rejected_leads_are_not_formal_results(collection_client: TestClient) -> None:
    collection_client.app.state.collection_service.verifiers = {"deepseek": RejectingVerifier()}
    response = collection_client.post(
        "/api/v1/collection-tasks",
        json={
            "name": "模型过滤",
            "request_text": "寻找软装公司",
            "cities": ["武汉市"],
            "keywords": ["软装公司"],
            "providers": ["web_search"],
            "models": ["deepseek"],
            "required_contacts": ["mobile"],
        },
    )
    task = response.json()

    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])
    before = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/leads").json()["items"]

    verification = collection_client.post(f"/api/v1/collection-tasks/{task['id']}/verify")
    after = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/leads").json()["items"]

    assert before[0]["verification_status"] == "unverified"
    assert verification.status_code == 202
    assert len(after) == 1
    assert after[0]["verification_status"] == "single_model_reject"
    assert after[0]["verification_models"] == ["deepseek"]


def test_partial_model_failure_retains_qualified_source_leads(collection_client: TestClient) -> None:
    collection_client.app.state.collection_service.verifiers = {
        "deepseek": FakeVerifier(),
        "doubao": FailingDoubaoVerifier(),
    }
    response = collection_client.post(
        "/api/v1/collection-tasks",
        json={
            "name": "双模型部分完成",
            "request_text": "寻找软装公司",
            "cities": ["武汉市"],
            "keywords": ["软装公司"],
            "providers": ["web_search"],
            "models": ["deepseek", "doubao"],
            "required_contacts": ["mobile"],
        },
    )
    task = response.json()

    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])
    before = collection_client.get(f"/api/v1/collection-tasks/{task['id']}").json()
    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/verify")

    leads = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/leads").json()["items"]
    status = collection_client.get(f"/api/v1/collection-tasks/{task['id']}").json()

    assert len(leads) == 1
    assert leads[0]["confidence_score"] >= 50
    assert before["verification_errors"] == 0
    assert status["verification_errors"] == 1
    assert status["verification_state"] == "completed_with_errors"


def test_json_export_omits_empty_cities_and_writes_unpublished(collection_client: TestClient) -> None:
    collection_client.app.state.collection_service.verifiers = {"deepseek": FakeVerifier()}
    task_response = collection_client.post(
        "/api/v1/collection-tasks",
        json={
            "name": "JSON 下载",
            "request_text": "寻找软装公司",
            "cities": ["武汉市", "长沙市"],
            "keywords": ["软装公司"],
            "providers": ["web_search"],
            "models": ["deepseek"],
            "required_contacts": ["mobile"],
        },
    )
    task = task_response.json()
    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])
    task_path = collection_client.app.state.collection_service.store.tasks_dir / f"{task['id']}.json"
    stored = __import__("json").loads(task_path.read_text(encoding="utf-8"))
    stored["leads"] = [lead for lead in stored["leads"] if lead["city"] == "武汉市"]
    collection_client.app.state.collection_service.store.save(stored)

    response = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/export?format=json")

    assert response.status_code == 200
    with ZipFile(BytesIO(response.content)) as archive:
        assert archive.namelist() == ["武汉市.json"]
        items = __import__("json").loads(archive.read("武汉市.json"))
        assert items[0]["email"] == "未公开"
        assert items[0]["wechat"] == "未公开"


def test_task_status_omits_large_internal_payloads(collection_client: TestClient) -> None:
    task = create_task(collection_client, ["武汉市"], ["软装公司"])
    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")

    response = collection_client.get(f"/api/v1/collection-tasks/{task['id']}")

    assert response.status_code == 200
    assert "leads" not in response.json()
    assert "activity" not in response.json()
    assert "completed_unit_keys" not in response.json()


def test_retry_errors_only_requeues_failed_units(collection_client: TestClient) -> None:
    collection_client.app.state.collection_service.providers = {
        "web_search": FailOnceWebSearchProvider()
    }
    task = create_task(collection_client, ["武汉市"], ["软装公司"])
    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    assert wait_for_task(collection_client, task["id"])["status"] == "completed_with_errors"

    retry = collection_client.post(f"/api/v1/collection-tasks/{task['id']}/retry-errors")

    assert retry.status_code == 202
    status_response = wait_for_task(collection_client, task["id"])
    assert status_response["status"] == "completed"
    assert status_response["lead_count"] == 1


def test_legacy_completed_zero_lead_task_can_retry_all_units(collection_client: TestClient) -> None:
    task = create_task(collection_client, ["武汉市"], ["软装公司"])
    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])
    service = collection_client.app.state.collection_service
    stored = service.get_task(task["id"])
    stored["leads"] = []
    stored["lead_count"] = 0
    stored["status"] = "completed"
    stored.pop("empty_units", None)
    service.store.save(stored)

    status_before = collection_client.get(f"/api/v1/collection-tasks/{task['id']}").json()
    retry = collection_client.post(f"/api/v1/collection-tasks/{task['id']}/retry-errors")
    status_after = wait_for_task(collection_client, task["id"])

    assert status_before["status"] == "completed_with_errors"
    assert retry.status_code == 202
    assert status_after["status"] == "completed"
    assert status_after["lead_count"] == 1


def test_delivery_pushes_real_leads_one_by_one_and_persists_status(collection_client: TestClient, monkeypatch) -> None:
    task = create_task(collection_client, ["武汉市", "长沙市"], ["软装公司"])
    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])
    received: list[dict[str, Any]] = []
    store = collection_client.app.state.collection_service.store
    original_iter_leads = store.iter_leads
    monkeypatch.setattr(store, "iter_leads", original_iter_leads)

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.headers["authorization"].startswith("Bearer ")
        received.append(__import__("json").loads(request.content))
        return httpx.Response(200, json={"errCode": 0, "errMsg": None})

    collection_client.app.state.delivery_client = httpx.Client(transport=httpx.MockTransport(handler))
    response = collection_client.post(
        f"/api/v1/collection-tasks/{task['id']}/deliver",
        json={
            "target_url": "https://receiver.example.test/leads",
            "bearer_token": token_urlsafe(24),
            "batch_size": 1,
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "completed"
    assert body["sent"] == 2
    assert body["failed"] == 0
    assert body["selected"] == 2
    assert [item["customerType"] for item in received] == ["B端客资", "B端客资"]
    assert [item["sourceAccountName"] for item in received] == ["豆包", "豆包"]
    persisted = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/leads").json()["items"]
    assert {lead["delivery_status"] for lead in persisted} == {"succeeded"}

    repeated = collection_client.post(
        f"/api/v1/collection-tasks/{task['id']}/deliver",
        json={"target_url": "https://receiver.example.test/leads", "batch_size": 1},
    )
    assert repeated.status_code == 200
    assert repeated.json()["sent"] == 0
    assert len(received) == 2


def test_batch_delivery_prefers_verified_then_rich_phone_fallback(collection_client: TestClient) -> None:
    task = create_task(collection_client, ["武汉市"], ["软装公司"])
    service = collection_client.app.state.collection_service
    stored = service.store.load(task["id"])
    base = {
        "city": "武汉市",
        "source_url": "https://example.test/company",
        "company_name_verification_status": "verified_public_web",
        "contact_verification_status": "verified_public_web",
        "contact_evidence_url": "https://example.test/contact",
        "verification_status": "single_model_match",
        "delivery_status": "display_only",
    }
    stored["leads"] = [
        {
            **base,
            "company_name": "武汉已核验邮箱公司",
            "email": "verified@example.test",
            "mobile": None,
            "service_intro": "软装设计",
        },
        {
            **base,
            "company_name": "武汉信息完整手机公司",
            "mobile": "13800138001",
            "email": "rich@example.test",
            "contact_person_name": "李明",
            "contact_person_role": "设计总监",
            "wechat": "liming-design",
            "address": "武汉市江汉区",
            "service_intro": "别墅软装与全案设计",
            "company_name_verification_status": "unverified",
            "contact_verification_status": "unverified",
        },
        {
            **base,
            "company_name": "武汉信息稀疏手机公司",
            "mobile": "13800138002",
            "email": None,
            "company_name_verification_status": "unverified",
            "contact_verification_status": "unverified",
        },
        {
            **base,
            "company_name": "武汉已经推送公司",
            "mobile": "13800138003",
            "delivery_status": "succeeded",
        },
    ]
    stored["lead_count"] = len(stored["leads"])
    service.store.save(stored)
    received: list[dict[str, Any]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        received.append(__import__("json").loads(request.content))
        return httpx.Response(200, json={"errCode": 0, "errMsg": None})

    collection_client.app.state.delivery_client = httpx.Client(transport=httpx.MockTransport(handler))
    response = collection_client.post(
        f"/api/v1/collection-tasks/{task['id']}/deliver",
        json={"target_url": "https://receiver.example.test/leads", "batch_size": 50},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["selected"] == 3
    assert body["sent"] == 3
    assert [payload["customerAccountName"] for payload in received] == [
        "武汉已核验邮箱公司",
        "武汉信息完整手机公司",
        "武汉信息稀疏手机公司",
    ]


def test_natural_language_task_uses_hidden_web_search_and_explicit_rules(collection_client: TestClient) -> None:
    service = collection_client.app.state.collection_service
    service.providers = {"web_search": FakeWebSearchProvider()}
    response = collection_client.post(
        "/api/v1/collection-tasks/from-natural-language",
        json={
            "request_text": "找武汉长沙软装公司，手机或邮箱，用 DeepSeek，只要 10 家",
            "models": ["deepseek"],
        },
    )

    assert response.status_code == 201
    task = response.json()
    assert task["cities"] == ["武汉市", "长沙市"]
    assert task["keywords"] == ["软装公司"]
    assert task["providers"] == ["web_search"]
    assert task["models"] == ["deepseek"]
    assert task["required_contacts"] == ["mobile", "email"]
    assert task["contact_match"] == "any"
    assert task["limit_per_city"] == 10


def test_natural_language_task_persists_continuous_verified_delivery_options(
    collection_client: TestClient,
) -> None:
    collection_client.app.state.erp_delivery.target_url = "https://erp.example.test/leads"

    response = collection_client.post(
        "/api/v1/collection-tasks/from-natural-language",
        json={
            "request_text": "查深圳别墅软装公司，需要手机号",
            "models": ["deepseek"],
            "continuous_mode": True,
            "verified_auto_deliver": True,
        },
    )

    assert response.status_code == 201
    task = response.json()
    assert task["continuous_mode"] is True
    assert task["verified_auto_deliver"] is True
    assert task["auto_deliver"] is False
    assert task["round_number"] == 1
    assert task["round_metrics"]["candidates"] == 0


def test_natural_language_task_rejects_two_automatic_delivery_modes(
    collection_client: TestClient,
) -> None:
    collection_client.app.state.erp_delivery.target_url = "https://erp.example.test/leads"

    response = collection_client.post(
        "/api/v1/collection-tasks/from-natural-language",
        json={
            "request_text": "查深圳别墅软装公司，需要手机号",
            "models": ["deepseek"],
            "auto_deliver": True,
            "verified_auto_deliver": True,
        },
    )

    assert response.status_code == 422


def test_verified_auto_delivery_requires_erp_configuration(
    collection_client: TestClient,
) -> None:
    response = collection_client.post(
        "/api/v1/collection-tasks/from-natural-language",
        json={
            "request_text": "查深圳别墅软装公司，需要手机号",
            "models": ["deepseek"],
            "verified_auto_deliver": True,
        },
    )

    assert response.status_code == 422


def test_natural_language_task_applies_defaults(collection_client: TestClient, monkeypatch) -> None:
    service = collection_client.app.state.collection_service
    service.providers = {"web_search": FakeWebSearchProvider()}
    service.verifiers = {"doubao": FakeVerifier()}
    service.verifiers["doubao"].id = "doubao"
    monkeypatch.setattr("app.api.routes.collection_tasks.tier3_plus_city_names", lambda: ["武汉市", "长沙市"])

    response = collection_client.post(
        "/api/v1/collection-tasks/from-natural-language",
        json={"request_text": "寻找全案软装公司", "models": ["doubao"]},
    )

    assert response.status_code == 201
    task = response.json()
    assert task["cities"] == ["武汉市", "长沙市"]
    assert task["models"] == ["doubao"]
    assert task["required_contacts"] == ["mobile"]
    assert task["contact_match"] == "all"


def test_natural_language_any_contact_and_score_50_are_retained(collection_client: TestClient) -> None:
    service = collection_client.app.state.collection_service
    service.providers = {"web_search": FakeWebSearchProvider()}
    response = collection_client.post(
        "/api/v1/collection-tasks/from-natural-language",
        json={"request_text": "找武汉软装公司，手机或邮箱，用 DeepSeek", "models": ["deepseek"]},
    )
    task = response.json()

    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])
    leads = collection_client.get(f"/api/v1/collection-tasks/{task['id']}/leads").json()["items"]

    assert len(leads) == 1
    assert leads[0]["mobile"]
    assert leads[0]["email"] is None
    assert leads[0]["confidence_score"] >= 50


def test_natural_language_task_uses_explicit_multi_model_selection(collection_client: TestClient) -> None:
    service = collection_client.app.state.collection_service
    doubao = FakeVerifier()
    doubao.id = "doubao"
    service.verifiers = {"deepseek": FakeVerifier(), "doubao": doubao}

    response = collection_client.post(
        "/api/v1/collection-tasks/from-natural-language",
        json={
            "request_text": "找武汉软装公司，使用 DeepSeek",
            "models": ["doubao", "deepseek"],
        },
    )

    assert response.status_code == 201
    assert response.json()["models"] == ["doubao", "deepseek"]


def test_existing_candidates_can_be_reenriched_by_doubao_without_rediscovery(
    collection_client: TestClient,
) -> None:
    service = collection_client.app.state.collection_service
    deepseek = ReenrichmentVerifier("deepseek")
    doubao = ReenrichmentVerifier("doubao")
    service.providers = {"web_search": EmptyWebSearchProvider()}
    service.verifiers = {"deepseek": deepseek, "doubao": doubao}
    task = create_task(collection_client, ["武汉市"], ["软装公司"])
    collection_client.post(f"/api/v1/collection-tasks/{task['id']}/start")
    wait_for_task(collection_client, task["id"])

    historical = service.get_task(task["id"])
    historical["leads"][0]["enrichment_model"] = "deepseek"
    historical["leads"][0]["mobile"] = None
    service.store.save(historical)

    response = collection_client.post(
        f"/api/v1/collection-tasks/{task['id']}/enrich-contacts"
    )
    wait_for_task(collection_client, task["id"])
    leads = collection_client.get(
        f"/api/v1/collection-tasks/{task['id']}/leads"
    ).json()["items"]

    assert response.status_code == 202
    assert deepseek.discovery_calls == 1
    assert deepseek.enrichment_calls == 0
    assert doubao.enrichment_calls == 2
    assert leads[0]["enrichment_model"] == "doubao"
    assert leads[0]["mobile"] == "139" + "12345678"


def test_delivery_status_and_retry_routes(collection_client: TestClient) -> None:
    task = create_task(collection_client, ["广州市"], ["软装公司"])

    status_response = collection_client.get(
        f"/api/v1/collection-tasks/{task['id']}/delivery-status"
    )
    retry_response = collection_client.post(
        f"/api/v1/collection-tasks/{task['id']}/retry-delivery"
    )

    assert status_response.status_code == 200
    assert status_response.json() == {"configured": False, "pending": 0}
    assert retry_response.status_code == 200
    assert retry_response.json() == {
        "attempted": 0,
        "succeeded": 0,
        "failed": 0,
    }


def test_manual_delivery_updates_one_lead_and_uses_erp_payload(
    collection_client: TestClient,
) -> None:
    task = create_task(collection_client, ["武汉市"], ["软装公司"])
    service = collection_client.app.state.collection_service
    stored = service.store.load(task["id"])
    stored["leads"] = [
        {
            "city": "武汉市",
            "company_name": "武汉示例软装有限公司",
            "normalized_company_name": "武汉示例软装有限公司".casefold(),
            "mobile": "13800138000",
            "service_intro": "别墅软装与全案设计",
            "saved_at": "2026-08-01 10:00:00",
            "delivery_status": "display_only",
        }
    ]
    stored["lead_count"] = 1
    service.store.save(stored)
    requests: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requests.append(request)
        return httpx.Response(200, json={"errCode": 0, "errMsg": None})

    collection_client.app.state.delivery_client = httpx.Client(
        transport=httpx.MockTransport(handler)
    )

    response = collection_client.post(
        f"/api/v1/collection-tasks/{task['id']}/leads/deliver",
        json={
            "city": "武汉市",
            "company_name": "武汉示例软装有限公司",
            "target_url": "https://erp.example.test/leads",
            "bearer_token": "secret-token",
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "succeeded"
    assert body["lead"]["delivery_status"] == "succeeded"
    assert requests[0].headers["authorization"] == "Bearer secret-token"
    payload = json.loads(requests[0].content)
    assert payload["customerAccountName"] == "武汉示例软装有限公司"
    assert payload["customerType"] == "B端客资"


def test_manual_delivery_accepts_email_without_mobile(
    collection_client: TestClient,
) -> None:
    task = create_task(collection_client, ["武汉市"], ["软装公司"])
    service = collection_client.app.state.collection_service
    stored = service.store.load(task["id"])
    stored["leads"] = [
        {
            "city": "武汉市",
            "company_name": "武汉无手机软装有限公司",
            "normalized_company_name": "武汉无手机软装有限公司".casefold(),
            "mobile": None,
            "email": "sales@example.test",
            "delivery_status": "display_only",
        }
    ]
    stored["lead_count"] = 1
    service.store.save(stored)

    requests: list[httpx.Request] = []
    def handler(request: httpx.Request) -> httpx.Response:
        requests.append(request)
        return httpx.Response(200, json={"errCode": 0, "errMsg": None})
    collection_client.app.state.delivery_client = httpx.Client(
        transport=httpx.MockTransport(handler)
    )

    response = collection_client.post(
        f"/api/v1/collection-tasks/{task['id']}/leads/deliver",
        json={
            "city": "武汉市",
            "company_name": "武汉无手机软装有限公司",
            "target_url": "https://erp.example.test/leads",
        },
    )

    assert response.status_code == 200
    assert response.json()["status"] == "succeeded"
    payload = json.loads(requests[0].content)
    assert payload["contactInfo"]["email"] == "sales@example.test"
    assert "mobilePhone" not in payload["contactInfo"]


def test_legacy_lead_service_intro_is_cleaned_when_results_are_read(
    collection_client: TestClient,
) -> None:
    task = create_task(collection_client, ["广州市"], ["软装公司"])
    service = collection_client.app.state.collection_service
    stored = service.store.load(task["id"])
    stored["leads"] = [
        {
            "city": "广州市",
            "company_name": "广州示例软装有限公司",
            "mobile": "13800138000",
            "service_intro": (
                "更新时间：2026-07-30。驾车导航，地址：广州市天河区示例路。"
                "专注别墅软装与全案设计。联系电话：13800138000。"
            ),
        }
    ]
    stored["lead_count"] = 1
    service.store.save(stored)

    response = collection_client.get(
        f"/api/v1/collection-tasks/{task['id']}/leads"
    )

    assert response.status_code == 200
    intro = response.json()["items"][0]["service_intro"]
    assert intro == "专注别墅软装与全案设计"
